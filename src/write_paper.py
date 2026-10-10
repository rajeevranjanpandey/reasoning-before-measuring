"""Build an evidence-linked manuscript; incomplete experiments remain explicit."""
import json
import re
import datetime
import subprocess
import shutil
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parents[1]

TITLE='Reasoning Before Measuring: Evaluating Small Language Models for Industrial Sensor Acquisition'
POLICY_LABELS={'fixed':'Ascending sensor ID','mi':'Mutual information','importance':'Tree importance',
               'forward':'Forward selection','random':'Random','tool':'Entropy tool','full':'All sensors',
               'qwen_fixed':'Qwen fixed plan','granite_fixed':'Granite fixed plan',
               'qwen_raw':'Qwen raw','qwen_assisted':'Qwen assisted',
               'granite_raw':'Granite raw','granite_assisted':'Granite assisted'}

def main():
    if (ROOT/'paper/manuscript.md').exists():
        raise RuntimeError('A complete-study manuscript exists. Use build_manuscript_docx.py for that source. Preserve the pilot record instead of overwriting the full paper.')
    analysis=json.loads((ROOT/'results/analysis.json').read_text())
    training=json.loads((ROOT/'results/training_summary.json').read_text())
    complete=analysis['primary_pilot_complete']
    audit_path=ROOT/'results/pilot_audit.json'
    audit=json.loads(audit_path.read_text()) if audit_path.exists() else None
    sensitivity_path=ROOT/'results/predictor_sensitivity.json'
    sensitivity=json.loads(sensitivity_path.read_text()) if sensitivity_path.exists() else None
    interface_path=ROOT/'results/interface_control.json'
    interface=json.loads(interface_path.read_text()) if interface_path.exists() else None
    lookup={(r['policy'],r['budget']):r for r in analysis['summaries']}
    def pct(policy,budget):return f"{lookup[(policy,budget)]['balanced_accuracy']*100:.2f}%"
    doc=Document();sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Times New Roman';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        if style.name in ['Title','Heading 1','Heading 2','Normal']:
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                fonts.attrib.pop(qn('w:'+key),None)
            fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
    doc.styles['Title'].font.size=Pt(19)
    doc.styles['Heading 1'].font.size=Pt(13)
    doc.styles['Heading 2'].font.size=Pt(11)
    markdown=[]
    def p(text):
        if not re.match(r'^\[\d+\]',text):
            text=text.replace('Delta(B;', 'Delta(B<SEMICOLON>')
            if text.startswith('Keywords:'):text=text.replace(';', ',')
            else:text=re.sub(r';\s*([A-Za-z])',lambda m:'. '+m.group(1).upper(),text)
            text=text.replace('<SEMICOLON>', ';')
        if re.match(r'^Figure \d+\.', text):
            text += (' Diagram code was prepared with ChatGPT/Codex. This schematic represents the documented protocol, not experimental outcomes.' if text.startswith(('Figure 1.', 'Figure 2.')) else ' Visualization code was prepared with ChatGPT/Codex. All plotted values derive from recorded experimental results.')
        paragraph=doc.add_paragraph()
        if re.match(r'^Figure \d+\.', text):
            paragraph.paragraph_format.keep_together=True
        parts=text.split('x_S')
        for i,part in enumerate(parts):
            if i:
                run=paragraph.add_run('x');run.italic=True
                run=paragraph.add_run('S');run.font.subscript=True
            paragraph.add_run(part)
        markdown.append(text+'\n')
    def h(text):doc.add_heading(text,level=1);markdown.append('## '+text+'\n')
    def sub(text):doc.add_heading(text,level=2);markdown.append('### '+text+'\n')
    def table(headers,rows,caption):
        p(caption);doc.paragraphs[-1].paragraph_format.keep_with_next=True
        t=doc.add_table(rows=1,cols=len(headers));t.style='Normal Table'
        t.autofit=False
        weights=[1.8,.4,1.3,1.5] if headers[1]=='B' else [1]*len(headers)
        widths=[Inches(6.57*w/sum(weights)) for w in weights]
        for column,width in zip(t.columns,widths):column.width=width
        for c,s in zip(t.rows[0].cells,headers):c.text=s
        trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single' if edge in ('top','bottom','insideH') else 'nil');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        t._tbl.tblPr.append(borders)
        for cell in t.rows[0].cells:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            cells=t.add_row().cells
            for cell,text in zip(cells,row):cell.text=str(text)
        for row in t.rows:
            prop=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');prop.append(keep)
            for column_index,cell in enumerate(row.cells):
                cell.width=widths[column_index]
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(3)
                    for run in para.runs:
                        run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
        markdown.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+
                        '\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
    doc.add_paragraph(TITLE,style='Title');markdown.append('# '+TITLE+'\n')
    p('A paired evaluation of acquisition quality and computational cost on Tennessee Eastman simulations')
    if not complete:
        p('Evidence status: this is an incomplete research manuscript. The classical pilot is complete; the full paired language-model pilot is not complete. No claim that language-model agents improve diagnosis or justify their cost is established.')
    h('Abstract')
    abstract=(
        'Repeated language-model consultation adds computation to sensor acquisition, but valid requests alone do not establish diagnostic value. '
        'We evaluate its marginal value in retrospective industrial fault diagnosis using a shared masked-input classifier on Tennessee Eastman Process simulations. '
        'The contribution is a paired empirical evaluation that separates acquisition behavior from diagnostic prediction and compares repeated consultation with static selection, an adaptive entropy tool and a one-time model plan. '
        'The pilot comprises 420 snapshots in 21 classes and 20 simulation-run families, with 5,040 model episodes across three decoding seeds. '
        f'At three sensors, assisted Qwen and Granite achieve {pct("qwen_assisted",3)} and {pct("granite_assisted",3)} balanced accuracy, versus {pct("forward",3)} for validation-selected forward acquisition. '
        'Their paired differences from forward selection are −8.81 and −13.65 percentage points; simultaneous bootstrap intervals exclude zero. '
        'Qwen repeats one acquisition trajectory, while Granite varies its choices without outperforming the forward reference. '
        'Recorded model runtime is higher, although measurement-scope differences preclude exact speedup estimates. '
        'Rescoring fixed acquisition paths with another classifier reverses the ordering of two classical policies. '
        'The results identify no marginal benefit over the selected classical reference in this configuration and show why acquisition validity, adaptation and diagnostic utility require separate assessment. '
        'The exploratory interpretation does not establish causal mechanisms, independent confirmation or general model equivalence.'
    )
    assert len(abstract.split())<=250,len(abstract.split())
    p(abstract)
    p('Keywords: active feature acquisition; industrial fault diagnosis; small language model; Tennessee Eastman Process; computational cost; reproducibility')
    h('1 Introduction')
    p('Here, industrial sensor acquisition means sequentially selecting which already-recorded sensor values to reveal to a diagnostic predictor under a fixed measurement budget. The evaluation simulates access to measurements; it does not acquire physical sensors or trigger new measurements in a plant. “Reasoning Before Measuring” names this selection-before-revelation task, rather than claiming that the experiment verifies a reasoning mechanism.')
    p('Industrial diagnostic services may have restricted access to measurements even when sensors already exist. A telemetry interface with a limited channel allowance is a motivating use case; no deployed interface or physical acquisition cost is measured in this study. Under this constraint, a useful selector must decide which information to request next from evidence already observed. Adding a language model creates a second resource cost: computation expended on each selection. An acquisition policy therefore needs to be assessed jointly by diagnostic quality at a stated sensor budget and by its measured computational overhead.')
    p('The central question is whether repeated consultation of a small language model (SLM) improves diagnosis enough to justify its additional computation once a diagnostic predictor and an acquisition tool are already available. We hold the classifier fixed across policies and change the selector. This design measures the effect of the implemented acquisition policy, conditional on that predictor, rather than conflating acquisition with language-model diagnostic capability.')
    p('The comparison separates three sources of apparent agent performance. A constrained interface can make requests valid, a policy can change its acquisition path across observations, and the resulting measurements can improve diagnosis. These properties need not coincide. A one-time plan provides a reference without repeated invocation, while a tool-only policy provides a reference for the numerical information available to the assisted selector. The original primary contrasts use validation-selected forward acquisition and the same-model fixed plan; the tool-only contrast is reported separately as descriptive analysis.')
    p('This study contributes a fully logged paired evaluation on common simulation snapshots, an analysis of acquisition concentration and tool-score agreement, and a predictor replay check that exposes dependence of policy rankings on the scoring classifier. These are empirical contributions within the tested setting. The study does not introduce a new acquisition algorithm or claim to be the first general benchmark of active feature acquisition.')
    p('The original protocol specified an improvement of at least three percentage points in balanced accuracy at three sensors, in both model families and against both primary references. That criterion is not met. The exploratory interpretation was developed after inspection of the pilot and retains the original unfavorable decision. We distinguish the prespecified comparison from subsequent behavioral observations and do not interpret identical predictions or nonsignificance as population equivalence.')
    p('Figure 1 shows the acquisition loop. The selector determines which information becomes available, while the shared classifier determines the diagnostic prediction. This separation allows acquisition policies to be compared on the same observations and scoring model. An adaptive interface permits a policy to respond to revealed values; it does not establish that the policy actually uses them.')
    doc.add_picture(str(ROOT/'paper/acquisition_workflow.png'),width=Inches(6.5))
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    markdown.append('![Acquisition loop](acquisition_workflow.png)\n')
    p('Figure 1. Sensor acquisition as controlled access to recorded measurements. Each episode begins without revealed values. A policy requests one previously unobserved sensor at a time; the common masked classifier updates its probabilities after each revelation. The loop stops at budget B. Only the evaluator accesses the ground-truth class, and no request changes the physical process. The primary repeated-model evaluation uses B = 3.')
    h('2 Related work and positioning')
    p('Feature selection can improve prediction, reduce computation and help interpret a dataset [9]. Static selection fixes a common subset, whereas sequential acquisition can choose the next measurement from values already observed. This distinction matters here because a language model can incur repeated computation even when its acquisition path stays fixed.')
    p('Active feature acquisition predates language-model agents. EDDI combines a partial variational autoencoder with expected information gain [3], while generative-surrogate acquisition models dependencies needed to value unrevealed features [4]. Covert et al. learn an amortized greedy acquisition policy by maximizing conditional mutual information [5]. These approaches establish that adaptive acquisition, information-based scoring, and inexpensive online policies are existing methodological families.')
    p('Shim et al. jointly learn a feature-acquisition agent and classifier, using a variable-size set representation for observed features [10]. Janisch et al. formulate requests and classification decisions as actions learned through deep reinforcement learning [11]. These learned policies are relevant comparators for a later evaluation, but we have not reproduced them. Our shared, fixed predictor isolates acquisition behavior within the implemented policies and answers a narrower question.')
    p('Huang et al. combine active feature querying with supervised matrix completion to reduce acquisition cost when constructing a useful training dataset [12]. Their setting concerns which missing feature values to acquire across the dataset. Our evaluation instead limits measurement access for each held-out snapshot after training; performance numbers from those settings are therefore not directly comparable.')
    p('Data-driven fault diagnosis also has a substantial industrial literature. Russell et al. describe principal component analysis, Fisher discriminant analysis, partial least squares and canonical variate analysis, with applications to the Tennessee Eastman simulator [13]. We use the simulator as a controlled source of labeled measurements, but do not reproduce that process-monitoring comparison or evaluate online detection delay.')
    p('Tool use by language models supplies a separate motivation. ReAct interleaves reasoning traces with actions that obtain information from an external environment [14], while Toolformer learns when and how to invoke application programming interfaces [15]. Our models receive a predefined sensor interface and, in the assisted condition, numerical acquisition scores. We do not train tool use or reproduce either method. A valid sensor request alone cannot establish useful information acquisition.')
    p('Perets et al. separate language-model knowledge elicitation from formal acquisition planning [6]. Their framework elicits discriminative statistics for a probabilistic model and considers binary classification and top-k identification. Our one-time sensor ordering is a control for repeated invocation; it is not a reproduction of their method. Likewise, the nearest-donor entropy tool used here is a transparent local baseline, not an implementation of EDDI or an estimator claimed to recover true conditional mutual information.')
    p('AFABench supplies a modular evaluation framework spanning static, greedy and reinforcement-learning acquisition methods, including a synthetic test of lookahead [21]. A general controlled benchmark is therefore not our novelty claim. Our narrower focus is the incremental diagnostic value and recorded computation of repeated local language-model consultation under a common industrial predictor. Cooper et al. use a greedy tree-search policy to balance anticipated information gain and action cost in a clinical diagnosis simulation [22]. That work motivates an explicit cost–information comparison, but its task and policy differ from our sensor selector; no direct numerical comparison is made.')
    p('The comparisons establish performance only against the implemented baselines. They do not establish superiority to published learned, amortized, or non-myopic acquisition methods. The scope is a reproducible benchmark of selected local acquisition policies under a common predictor, rather than a new acquisition algorithm or a state-of-the-art claim.')
    h('3 Problem formulation and evaluation targets')
    p('Let x contain 41 measured process variables and let y denote one of 21 operating classes. An episode begins with an empty observed set S. At each step, the policy receives the acquired values x_S, the legal sensor identities, remaining budget, fixed sensor descriptions, and the shared classifier probabilities p(y | x_S, S). It selects one previously unobserved sensor. The evaluator then reveals that sensor value. The policy never receives the full held-out row, its ground-truth label, simulation-run identifier, or sample index.')
    p('An episode terminates after B successful acquisitions, where B is 1, 3, or 5. The primary endpoint uses B = 3. All policies are assessed with the same final classifier. The full-sensor classifier is a reference condition rather than a mathematical upper bound: finite-sample classification performance need not be monotone in the number of features. No free initial sensor pack is supplied.')
    p('The repeated language-model policies are evaluated only at the primary budget of three sensors. Classical and fixed-plan policies additionally supply one- and five-sensor results. Consequently, their budget curves do not represent independently executed repeated-model trials at one or five sensors.')
    p('The setting is frozen retrospective replay. A measurement does not change while an agent reasons, and requesting it has no physical effect on the process. Recorded composition analyzers can also represent delayed measurements in the simulator. Consequently, the study does not measure online detection delay, actuator decisions, plant safety, energy saved by switching off sensors, or the consequence of stale telemetry.')
    p('For a fixed predictor q, the estimand at budget B is Delta(B; pi, pi0, q) = BA(pi, q, B) minus BA(pi0, q, B), where BA denotes balanced accuracy and pi and pi0 denote acquisition policies. The comparison is conditional on the fitted predictor, interface, sampled decoding seeds and snapshot distribution. It measures the complete policy configuration; the assisted-versus-fixed-plan contrast does not isolate repeated invocation from access to observations and tool scores.')
    h('4 Dataset and leakage controls')
    p('The Tennessee Eastman Process is a simulated plant-wide control benchmark introduced by Downs and Vogel [2]. The present evaluation uses a subsequent public simulation dataset rather than measurements from an operating industrial plant.')
    p('We use version 1.0 of Additional Tennessee Eastman Process Simulation Data for Anomaly Detection Evaluation, deposited by Rieth et al. [1]. The authoritative metadata describes 500 simulations for each of 20 faults and normal operation, with distinct random seeds between the publisher training and testing partitions. Training files contain 500 samples per run, and testing files contain 960 samples, sampled every three minutes. Faults begin after one hour in training and eight hours in testing. All four downloaded files match the publisher MD5 checksums; local SHA-256 hashes are retained.')
    p('One snapshot at sample 320 is selected from every run before outcome evaluation. At a three-minute sampling interval, sample 320 is approximately 16 hours into a run, approximately 15 hours after training fault onset and eight hours after testing fault onset; a zero-origin indexing convention would shift these times by one sampling interval. The mismatch in post-onset duration is retained and reported as a limitation rather than silently aligning or selecting favorable time points. Normal runs use the same elapsed time. The candidate set comprises xmeas_1 through xmeas_41; the 11 manipulated-variable columns and three metadata columns are excluded.')
    p('Within the publisher training partition, runs 1–350 supply 7,350 training snapshots, runs 351–400 supply 1,050 validation snapshots, and runs 401–420 supply the development pool. Run identifiers are grouped across all fault classes to avoid separating potentially shared seed families. The scientific pilot uses publisher testing runs 1–20, yielding 420 snapshots. Testing runs 21–70 are allocated to a separately frozen confirmation, and runs 71–90 to a subsequent predictor study. Neither supplies results reported in this pilot manuscript. Runs 91–500 remain reserved. Snapshot conversion alone does not constitute outcome evaluation.')
    p('A memory-bounded reader processes the pinned R serialization one numeric column at a time. It rejects unexpected headers, numeric types, dimensions, duplicate run-class identities, and column names. On the complete normal training file, its 27,500 retained snapshot cells and column order agree exactly with the independent pyreadr reader. This engineering check supports extraction correctness; it is not a diagnostic-performance result.')
    p('Figure 2 distinguishes model fitting, policy selection, interface development and outcome evaluation. Run-family grouping reduces overlap between these roles; it does not remove the publisher partitions\' distribution differences. The separate confirmation and predictor-extension partitions are shown to clarify their exclusion from the results reported here.')
    doc.add_picture(str(ROOT/'paper/study_design.png'),width=Inches(6.5))
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    markdown.append('![Study partitions](study_design.png)\n')
    p('Figure 2. Partition roles in the study. All reported pilot outcomes use 420 snapshots from testing runs 1–20, comprising one snapshot for each of 21 operating classes in each run family. Training, validation and development use the publisher training partition. The development pool is not a second test set. Confirmation and predictor-extension allocations are separate studies and contribute no results to this manuscript. Diagram areas are not proportional to sample counts.')
    h('5 Shared predictor and acquisition policies')
    p('We use an Extra Trees classifier [16], implemented in scikit-learn [17], with 64 trees, maximum depth 14, and a minimum of three samples per leaf. Training means and standard deviations standardize each channel. Each input concatenates 41 masked standardized values and 41 explicit observation indicators. Sixteen mask assignments per training snapshot cycle through budgets zero to five. Nonempty masks are sampled independently, zero-budget masks are repeated deterministic empty sets, and one full-observation example is added. A zero in a hidden value position is therefore distinguishable from an observed standardized zero. The classifier, scaler and training-only donor bank are shared by all policies.')
    p('The classical suite includes an ascending sensor-ID reference order, a training-only marginal mutual-information ranking, a training-only tree-importance ranking, a validation-selected forward order, a seeded random policy, and an adaptive expected-entropy tool. Forward selection adds the sensor with the highest validation balanced accuracy at each of five steps. The classical champion is selected on validation accuracy at three sensors before pilot evaluation. In this run, it is forward selection. Its first three zero-based feature indices are 37, 17 and 20, corresponding to sensor IDs 38, 18 and 21.')
    p('For a candidate j, the tool estimates H(p(y | x_S)) minus the mean entropy after revealing j under eight donor values. Donors are the nearest training snapshots using distances restricted to acquired standardized channels. At the empty state, eight training donors are selected using a fixed random seed. The actual hidden test value is never used to choose donors or score candidates. Conditional approximation error and classifier miscalibration can affect these scores. Tool-assisted prompts serialize scores to five decimal places; the tool-only baseline uses numerical scores before text serialization. This precision difference is an explicit implementation limitation.')
    p('The repeated raw selector sees the permitted state without entropy scores. The repeated tool-assisted selector additionally sees the score table. Both use the same shared diagnosis probabilities. A separate fixed-plan control requests five unique sensors from the same model once, using the same dictionary and empty-state diagnosis probabilities. It receives no additional training ranking or held-out values. The plan is reused across episodes and is not presented as a reproduction of prior probabilistic elicitation work. The tool-assisted contrast with this plan includes both adaptive evidence access and tool access; the separate tool-only contrast is needed to assess the incremental value of the SLM.')
    p('Probability quality matters because the entropy tool scores candidate measurements using the shared classifier probabilities. Niculescu-Mizil and Caruana examine probability estimation across supervised learners [18]; Guo et al. examine calibration in modern neural networks [19]. These studies motivate checking calibration separately from classification accuracy. We have not performed a dedicated calibration experiment, so the present results cannot attribute tool behavior to calibrated or miscalibrated probabilities.')
    p('The entropy score is a heuristic. Its interpretation as conditional mutual information would require coherent conditional probabilities and donor values representing the relevant conditional distribution, neither of which is established here. Scores may be negative, and increased classifier confidence need not improve diagnostic correctness. No greedy approximation guarantee is claimed. At budget three, 10,660 sensor subsets are possible; training supplies three randomly sampled three-sensor masks per snapshot. Performance on policy-selected masks therefore requires empirical validation.')
    h('6 Language models and constrained action interface')
    p('The model families are Qwen3-1.7B and IBM Granite 3.3 2B Instruct [7,8]. Official Q8_0 GGUF artifacts are pinned to repository revisions and local SHA-256 hashes. Each model is loaded alone using llama-cpp-python 0.3.19 with Metal acceleration, four CPU threads and a 4,096-token context. The verified machine is an Apple M2 Mac mini with 8 GiB of unified memory running macOS 26.6.2. This differs from the 16 GB assumption in the original planning document. A preliminary CPU smoke test is excluded from reported experimental latency.')
    p('The model-provided chat templates are retained. Thinking is disabled through template variables; the Granite date function is fixed to the experiment date for reproducibility. Action generation uses temperature 0.7, top-p 0.8, top-k 20, a 32-token cap, and a 90-second process deadline. The scientific pilot specifies decoding seeds 17, 29 and 43, with deterministic offsets for action and repair attempts. The same settings do not imply identical computational effort across tokenizers or model families.')
    p('Initial unconstrained development traces revealed bare-integer and prose-wrapped outputs from Granite. A shape-only JSON grammar exposed repeated requests for an already acquired sensor. Both sets of traces are preserved. A second documented engineering amendment applies the same JSON grammar enumerating legal remaining sensor IDs to both model families and repeats the development check. The grammar receives no hidden values or labels. Legal-action success is enforced by the interface and cannot be credited to autonomous model reliability. The primary experiment therefore evaluates legal-action-constrained acquisition selectors; the earlier failure traces describe separate interface conditions.')
    p('Each invalid action receives one repair attempt. After a second invalid response, the fallback selects the first legal sensor in the fixed training marginal-information ranking. The fallback is recorded, consumes no extra sensor budget, and contributes its time and resulting outcome to the main denominator. A hard timeout terminates the model process. No failed episode is silently removed. The development criterion requires at least 19 of 20 trajectories in every model–policy combination to finish without fallback; development identities are shared across cells and are not counted as independent replicated datasets.')
    h('7 Statistical analysis and resource accounting')
    p('Balanced accuracy is the arithmetic mean of class-specific recalls. For stochastic policies, correctness is first averaged across decoding seeds within each run-class identity. We use bootstrap resampling [20] to estimate uncertainty, drawing 20,000 paired resamples of simulation-run families and retaining all 21 classes within each resampled family. Resampling whole families is our design choice to preserve the dependence structure assumed in this dataset. Thus 420 snapshots correspond to 20 resampling clusters, not 420 independent clusters, and three decoding seeds do not triple the sample size.')
    p('The four primary contrasts compare each tool-assisted model with the frozen classical champion and with its same-model fixed plan. The planned simultaneous intervals use 98.75% marginal bootstrap intervals, corresponding to a Bonferroni allocation across four contrasts. Descriptive policy intervals use 95% coverage. A positive pilot requires a difference of at least three percentage points and a positive simultaneous lower bound for every contrast, with no reversal across decoding seeds. Approximate bootstrap coverage with only 20 clusters remains a limitation.')
    p('Model load time is separate from warmed episode time. Episode timing includes prompt construction, acquisition scoring where applicable, model calls, repairs, fallback handling and prediction. Every request renders a complete new state, while the runtime can reuse an identical token prefix after removing the previous differing suffix. No previous episode messages are included. Reported timings therefore describe warmed prefix reuse; earlier uncached development timings are archived separately. Process peak resident memory is reported with its measurement scope. It is not equated with total GPU or unified-memory usage. System swap observations include other applications and cannot be attributed exclusively to the experiment.')
    p('The Granite pilot was interrupted after 1,861 completed episodes and resumed for the remaining 659 identities on 21 September 2026. Scientific settings and completed records were retained. Runtime initialization files describe the latest load; per-episode timing and memory records survive the restart. Timing summaries therefore combine two execution sessions and are descriptive rather than controlled hardware microbenchmarks.')
    p('Classical timings at smaller budgets are checkpoints within a five-acquisition trajectory. They include earlier checkpoint predictions and therefore are not fresh, isolated budget-specific latency trials. These low-cost timings provide descriptive context but should not support precise speedup ratios. Model pretraining costs are excluded. Local classifier fitting and policy setup are retained separately; deployment amortization would divide setup time by an explicitly stated number of future episodes.')
    p('OpenAI Codex, a GPT-6-based assistant, supported research coding, analysis scripting and manuscript preparation; the application build was not recorded. Result figures are generated deterministically from recorded numerical results using Matplotlib. The two opening schematics depict the documented protocol and encode no measured outcomes; all figure-generation code is retained. No generative image model was used to create or alter research measurements. This assistance is separate from the Qwen and Granite systems evaluated in the experiment.')
    if sensitivity or interface:
        p('Two supplementary checks were specified during completion of the pilot and are treated as post-pilot descriptive analyses. An interface control asks each model to return sensor IDs 7, 23 and 41 while all 41 IDs remain legal. A predictor sensitivity check changes only the maximum tree depth from 14 to unlimited, retains the original training rows and masks, and rescores the recorded acquisition paths. It does not rerun acquisition under changed diagnosis probabilities or tool scores. These checks do not modify the primary decision rule.')
    p('All family-bootstrap intervals condition on the three sampled decoding seeds and the fitted predictor. They do not incorporate predictor-refitting or fixed-plan-generation uncertainty. Identical step-specific seeds are reused across episodes, creating common randomness rather than independent decoding for every snapshot. Per-seed results are therefore reported separately. A point estimate of at least three percentage points with a lower bound above zero establishes a positive effect whose estimate meets the target; it does not establish a population effect of at least three points, which would require the lower bound itself to exceed three points.')
    p('Energy consumption and carbon emissions were not measured. Latency and process memory are resource descriptors and are not estimates of environmental impact. Regeneration of tables from saved records, prediction replay from pinned artifacts, and complete model re-execution are distinct reproducibility tasks; cross-platform bitwise determinism is not assumed.')
    h('8 Results')
    sub('8.1 Diagnostic performance under a sensor budget')
    p(f'Fitting the shared classifier took {training["fit_seconds"]:.2f} seconds, and measured predictor fitting plus classical policy selection took {training["setup_seconds"]:.2f} seconds in the recorded run. Validation balanced accuracy at three sensors was {100*training["validation"]["forward"]:.2f}% for forward selection and {100*training["validation"]["tool"]:.2f}% for the adaptive tool. The forward order was frozen as the primary classical comparator.')
    table(['Policy','B','Balanced accuracy','95% interval'],[
        [POLICY_LABELS[r['policy']],r['budget'],f"{100*r['balanced_accuracy']:.2f}%",f"{100*r['ci95_low']:.2f}–{100*r['ci95_high']:.2f}%"]
        for r in analysis['summaries'] if r['budget'] in (3,41)],
        'Table 1. Pilot results at the primary budget and the full-sensor reference. Every row uses the same 420 snapshot identities; intervals resample 20 run families.')
    p(f'At three sensors, forward selection reaches {pct("forward",3)}, compared with {pct("tool",3)} for tool-only acquisition and {pct("random",3)} for random acquisition. At five sensors, tool-only acquisition reaches {pct("tool",5)} while forward selection reaches {pct("forward",5)}. The ordering of these two policies changes with budget. These observations do not identify the value of an SLM; they establish the classical reference conditions that an SLM must improve upon.')
    fig=ROOT/'paper/accuracy_by_budget.png'
    if fig.exists():
        doc.add_picture(str(fig),width=Inches(5.4));markdown.append('![Accuracy by budget](accuracy_by_budget.png)\n')
        doc.paragraphs[-1].alignment=1
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p('Figure 3. Balanced accuracy by sensor budget. Error bars are pointwise 95% bootstrap intervals from 20,000 resamples of 20 run families, retaining all 21 classes. Each policy uses the same 420 snapshots. Random-policy outcomes average three seeds within each identity. The dashed line and gray band show the 41-sensor reference and its 95% interval. Panels separate policy groups for readability and share the same vertical scale. Connecting lines link measured budgets only. Intervals are descriptive and are not simultaneous across policies or budgets.')
    if analysis['development']:
        table(['Model and variant','Finished','Fallback episodes','Median seconds'],[
            [r['model']+' '+r['variant'],r['n'],r['fallback_episodes'],f"{r['median_seconds']:.2f}"]
            for r in analysis['development']],
            'Table 2. Constrained development execution checks. Each cell contains 20 trajectories. These checks concern execution and fallback behavior, not comparative diagnostic accuracy.')
    if complete:
        sub('8.2 Marginal value of repeated consultation')
        table(['Model','Reference','Difference pp','Simultaneous interval pp'],[
            [r['model'].title(),POLICY_LABELS[r['comparator']],f"{100*r['difference']:.2f}",f"[{100*r['simultaneous_low']:.2f}, {100*r['simultaneous_high']:.2f}]"]
            for r in analysis['contrasts']],
            'Table 3. Prespecified paired primary contrasts at three sensors. Positive differences favor the tool-assisted language-model selector.')
        p('The decision under the locally frozen analysis plan is: '+analysis['decision']+'. The plan was not publicly preregistered. This decision concerns the specified models, predictor, prompt and snapshot distribution. It is not a universal assessment of small language models.')
        p(f'The raw Qwen and Granite policies attain {pct("qwen_raw",3)} and {pct("granite_raw",3)}, whereas the tool-assisted policies attain {pct("qwen_assisted",3)} and {pct("granite_assisted",3)}. Both one-time model plans select sensor IDs 1, 2, 3, 4 and 5 in that order, so their three-sensor outcomes coincide with the ascending-ID baseline. These fixed plans are single sampled plans, not averages over plan-generation seeds.')
        doc.add_picture(str(ROOT/'paper/paired_contrasts.png'),width=Inches(6.1))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p('Figure 4. Four primary paired contrasts. Error bars are 98.75% marginal run-family bootstrap intervals under the prespecified Bonferroni allocation. The dashed line marks the three-percentage-point effect target. Positive values favor the tool-assisted model.')
        markdown.append('![Primary paired contrasts](paired_contrasts.png)\n')
        if audit:
            seed_text=[]
            for model,item in audit['models'].items():
                for variant,cell in item['cells'].items():
                    vals=', '.join(f"{100*cell['accuracy_by_seed'][str(seed)]:.2f}%" for seed in (17,29,43))
                    seed_text.append(f'{model.title()} {variant}: {vals}')
            p('Balanced accuracy for decoding seeds 17, 29 and 43, respectively, is ' + '; '.join(seed_text) + '. These are repeated measurements on the same snapshots, not independent datasets. The zero-width Qwen-versus-fixed-plan interval records identical sampled outcomes and does not establish population equivalence.')
        sub('8.3 Recorded computation and timing scope')
        table(['Policy','Median seconds','95th percentile seconds'],[
            [POLICY_LABELS[r['policy']],f"{r['median_seconds']:.4f}",f"{r['p95_seconds']:.4f}"]
            for r in analysis['summaries'] if r['budget']==3 and r['policy'] in ('forward','tool','qwen_fixed','granite_fixed','qwen_raw','qwen_assisted','granite_raw','granite_assisted')],
            'Table 4. Recorded episode latency at three sensors. Classical checkpoints and model episode trials have different timing scopes, as described in Section 7; exact speedup ratios are not inferred.')
        plans={m:json.loads((ROOT/'results'/f'{m}_fixed_plan.json').read_text()) for m in ('qwen','granite')}
        p(f"The one-time fixed-plan requests take {plans['qwen']['setup_seconds']:.2f} seconds for Qwen and {plans['granite']['setup_seconds']:.2f} seconds for Granite, excluding model loading. Reusing these plans amortizes that request time across subsequent episodes. Shared predictor fitting and validation policy selection are reported separately above; neither source-data conversion nor pretrained-model development cost is included in online episode latency.")
        doc.add_picture(str(ROOT/'paper/accuracy_latency.png'),width=Inches(5.8))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p('Figure 5. Recorded runtime at three sensors. Circles show medians and diamonds show 95th percentiles; connecting segments are descriptive spans, not confidence intervals. Panel A shows 1,260 repeated-model episodes per policy in seconds. Panel B shows 420 reference records per policy in milliseconds, distinguishing classical trajectory checkpoints from fixed-plan replay. Axes use separate linear scales. Setup and model loading are excluded, and timing-scope differences preclude precise speedup ratios.')
        markdown.append('![Median and upper-tail recorded runtime](accuracy_latency.png)\n')
        if analysis.get('descriptive_contrasts'):
            table(['Assisted model','Comparator','Difference pp','95% interval pp'],[
                [r['model'].title(),POLICY_LABELS[r['comparator']],f"{100*r['difference']:.2f}",f"[{100*r['ci95_low']:.2f}, {100*r['ci95_high']:.2f}]"]
                for r in analysis['descriptive_contrasts']],
                'Table 5. Additional descriptive paired contrasts. These post-pilot intervals are not adjusted for multiplicity and do not replace the four primary contrasts.')
        sub('8.4 Acquisition behavior and interface controls')
        if audit and audit['complete']:
            p('The post-run audit verifies all 5,040 expected episode identities and 15,120 acquisitions, with no duplicate identities, hidden-value mismatch, illegal acquisition, or discrepancy between stored and recomputed final predictions. Prompt-state fields contain only permitted observations, legal identities, remaining budget, shared probabilities and, for assisted policies, the score table. This validates the recorded interfaces; it does not prove the absence of every possible implementation defect.')
            table(['Model policy','Unique trajectories','Fallback episodes','Tool agreement','Worker peak RSS GiB'],[
                [model+' '+variant,cell['distinct_trajectories'],cell['fallback_episodes'],
                 'Not applicable' if cell['rounded_tool_top_choice_agreement'] is None else f"{100*cell['rounded_tool_top_choice_agreement']:.1f}%",f"{cell['max_worker_rss_bytes']/2**30:.2f}"]
                for model,item in audit['models'].items() for variant,cell in item['cells'].items()],
                'Table 6. Descriptive behavior and recorded process memory across 1,260 episodes per cell. Tool agreement uses rounded prompt scores, with the lowest sensor ID breaking ties. Peak resident set size (RSS) is cumulative within a worker process, not per-policy memory demand or total unified-memory use. Legal choices are enforced by the grammar.')
            p('Trajectory frequencies and tool-choice agreement are descriptive mechanism diagnostics. They cannot establish that sensor semantics cause a choice; that would require a matched intervention such as anonymizing descriptions. Class-specific recalls, decoding-seed outcomes, token counts and measured process-memory peaks are supplied in the machine-readable audit.')
            p('Qwen selects the same ordered trajectory, sensor IDs 1, 2 and 3, in every raw and assisted episode across all three decoding seeds. Its repeated calls therefore do not adapt the observed acquisition path in this sample. Granite produces eight distinct raw trajectories and 32 assisted trajectories. Its assisted choices match the highest rounded tool score in only 0.82% of decisions, compared with 33.33% for Qwen. Agreement alone does not establish causal use of the score table.')
            if (ROOT/'paper/sensor_selection.png').exists():
                figure_info=json.loads((ROOT/'results/diagnostic_figures.json').read_text())
                doc.add_picture(str(ROOT/'paper/sensor_selection.png'),width=Inches(6.3))
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
                p(f"Figure 6. Acquisition concentration by model and policy. Values are the percentage of episodes selecting each sensor at any of the three steps. The 12 most frequently selected channels account for {100*figure_info['selected_sensor_coverage']:.2f}% of all recorded model acquisitions; omitted channels are not assumed to have zero use. Each row contains 1,260 episodes on the same 420 identities.")
                markdown.append('![Sensor selection concentration](sensor_selection.png)\n')
                doc.add_picture(str(ROOT/'paper/class_recall.png'),width=Inches(6.3))
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
                p('Figure 7. Class-specific recall on the same 20 run families. Colors show descriptive point estimates, averaging model decoding seeds within each class. Class 0 denotes normal operation; classes 1–20 follow the dataset fault numbering. The full-sensor row uses 41 channels, while all other rows use three. The two fixed model plans yield identical predictions and are shown once.')
                markdown.append('![Class-specific recall](class_recall.png)\n')
                if interface:
                    p(f"In the supplementary interface control, {interface['passed']} of {interface['total']} responses match the explicitly requested sensor ID. This checks whether the constrained interface can produce these legal IDs; it is not a fault-diagnosis accuracy test or evidence of unconstrained output reliability.")
                    if interface['passed']==6:
                        p('Both models can emit sensor IDs 7, 23 and 41 under explicit requests. Thus the grammar does not force the repeated Qwen trajectory to IDs 1, 2 and 3. The control does not distinguish among prompt design, model behavior and sampling effects as explanations for the observed acquisition concentration.')
                p('Figure 8 compares the assisted models with the two principal classical references across all 21 classes. Qwen assisted has zero observed recall in 12 classes, compared with six for Granite assisted, yet Qwen has higher balanced accuracy. Recovering a few cases in more classes is therefore distinct from improving average recall. Normal-operation recall is 0% for Qwen assisted and forward selection, 20% for Granite assisted and 5% for the entropy tool. All four policies have zero observed recall for fault classes 3, 5 and 7. These descriptive failure patterns argue against interpreting aggregate accuracy as broad diagnostic coverage. They do not establish population failure rates.')
                doc.add_picture(str(ROOT/'paper/class_recall_spider.png'),width=Inches(6.5))
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
                p('Figure 8. Spider charts of class-specific recall at three sensors. Panels A and B compare Qwen assisted and Granite assisted with the same forward and entropy-tool references. Every spoke uses the same linear 0–100% recall scale. N denotes normal operation and 1–20 follow the dataset fault IDs. Estimates use 20 run families, averaging the three model decoding seeds within each identity. All classes are retained in dataset order. Connecting lines aid visual comparison only, with no area-based score, class-adjacency interpretation or classwise confidence intervals. Figure 7 and the accompanying CSV provide complementary class-level detail.')
                markdown.append('![Spider comparison of class recall](class_recall_spider.png)\n')
                p('Supplementary Figures S1–S4 show the full confusion matrices for forward selection, the entropy tool and both assisted models on the same pilot identities. They distinguish missed normal-operation cases from confusion among fault classes. The matrices retain all 21 true and predicted classes, with exact percentages supplied as a CSV. They are descriptive analyses of the completed pilot, not additional test sets or cellwise significance tests.')

        if sensitivity:
            sub('8.5 Dependence on the scoring predictor')
            table(['Policy','Original accuracy','Replay accuracy','Replay 95% interval'],[
                [POLICY_LABELS[r['policy']],pct(r['policy'],r['budget']),f"{100*r['accuracy']:.2f}%",f"{100*r['ci95_low']:.2f}–{100*r['ci95_high']:.2f}%"]
                for r in sensitivity['summaries'] if r['policy'] in ('forward','tool','full','qwen_assisted','granite_assisted','qwen_fixed','granite_fixed')],
                'Table 7. Supplementary predictor replay sensitivity. The alternate classifier has unlimited tree depth. Acquisition paths are held fixed; intervals are descriptive and unadjusted.')
            p('Replay sensitivity measures how the same acquired measurements are scored by another predictor. It cannot show how an adaptive policy would behave if the alternate predictor had supplied its probability states during acquisition. The full set of replay predictions and paired contrasts is retained with the supplementary results.')
            alt={r['policy']:r['accuracy'] for r in sensitivity['summaries']}
            p(f"With unlimited tree depth, the full-sensor reference reaches {100*alt['full']:.2f}% and replay of the entropy-tool paths reaches {100*alt['tool']:.2f}%. Forward selection and the repeated Qwen paths both reach {100*alt['forward']:.2f}%. The ordering of forward selection and tool-only acquisition therefore depends on the scoring predictor. This limits robustness of the primary policy ranking and reinforces the need to separate the frozen primary result from post-pilot replay sensitivity.")
    else:
        p('The four primary language-model contrasts are not reported because the full paired model pilot is incomplete. Development success or failure cannot substitute for those comparisons. No language-model accuracy, superiority, equivalence, or confirmed compute–quality frontier is claimed from the current record.')
    h('9 Discussion and threats to validity')
    if complete:
        p(f'At three sensors, the tool-assisted models do not improve on the validation-selected classical policy. Their paired differences from forward selection are {100*analysis["contrasts"][0]["difference"]:.2f} percentage points for Qwen and {100*analysis["contrasts"][2]["difference"]:.2f} for Granite. The recorded model episode times are also greater than those of the reference policies, although the timing-scope differences prevent precise speedup ratios. The specified repeated consultations therefore provide no measured quality benefit that offsets their observed runtime in this configuration.')
    sub('9.1 What the comparisons establish')
    p('Acquisition diversity and diagnostic value should be interpreted separately. Qwen shows no observed trajectory adaptation in this sample, while Granite changes trajectories without surpassing the forward reference. The six requested-ID controls establish that the grammar can emit alternative IDs, but do not explain the repeated selections. A future randomized comparison of sensor ordering, anonymized descriptions and tool-score visibility could distinguish positional, semantic and score-following effects. These are proposed tests, not mechanisms established by the present study.')
    p('A single sensor budget is insufficient to characterize the classical policies: the validation-selected order is strongest among the recorded classical methods at three sensors, whereas the adaptive tool is stronger at five. This pattern motivates keeping sensor budget explicit instead of compressing quality and latency into an arbitrary scalar reward. Repeated-model performance at one and five sensors is not reported in this pilot.')
    sub('9.2 Predictor and distribution limitations')
    p(f'The full-sensor reference accuracy of {pct("full",41)} is a material limitation. A single late snapshot and a classifier trained across very different observation masks may provide a weak diagnostic foundation. The change in post-onset duration between training and testing can further affect generalization. These explanations are hypotheses, not established causal findings. The frozen primary predictor and the separately declared replay sensitivity answer different questions. Replay changes the scoring of observed paths, but does not test adaptation to a new predictor. Future classifier selection should use validation data, followed by closed-loop evaluation on untouched confirmation runs.')
    p('The donor tool uses only eight samples and an observed-coordinate nearest-neighbor approximation. It need not capture multimodal conditional distributions or long-range sensor dependencies. Its performance therefore cannot stand for the entire classical acquisition literature. Similarly, the SLMs receive sensor descriptions and numerical class probabilities but not a complete textual fault taxonomy or fine-tuning on industrial acquisition trajectories. Poor performance would bound this prompt and interface, not the capabilities of all possible industrial agents.')
    sub('9.3 Generalization and confirmatory scope')
    p('The study includes one simulator, one snapshot timestamp, one quantization level and one hardware configuration. Training and testing seed families are separated according to publisher documentation, but their actual generator seeds are unavailable. Only 20 test run families enter the pilot. The finite-cluster bootstrap, stochastic-decoding variability, system activity during timing, and the difference between process RSS and total unified-memory pressure constrain precision and portability. No real plant or unrelated industrial dataset is included.')
    if complete:
        p('This completed pilot supports conclusions about the tested configuration only. Independent confirmation, closed-loop re-evaluation with an alternate shared predictor, learned and non-myopic acquisition baselines, and randomized mechanism controls would be needed to support a broader positive or equivalence claim. These extensions were not executed in this pilot. The retained stop decision applies to progression of the specified positive-effect hypothesis. The exploratory observations instead identify testable questions about acquisition concentration, tool use and predictor dependence; they do not establish that all small language models lack acquisition value.')
    else:
        p('The paired model comparisons remain incomplete. Independent confirmation and broader comparisons must be considered after the prespecified pilot decision; the present partial record cannot resolve the research question.')
    sub('9.4 Coverage gaps')
    p('Table 8 identifies the claims that the present design can support and the questions it leaves unresolved. Uniform sensor-count budgets measure access to information, but do not represent heterogeneous monetary cost, measurement delay or energy use. A single late snapshot also excludes the transition from healthy operation to a developing fault. These boundaries are central to interpretation because a policy that performs poorly in retrospective diagnosis could behave differently when observations arrive over time or have unequal costs.')
    table(['Dimension','Covered in this study','Not established'],[
        ['Information access','Sequential revelation of recorded values','Physical sensing, streaming data or sensor control'],
        ['Diagnostic setting','One simulator, 21 classes, sample 320','External plants, early detection or distribution shift'],
        ['Policy comparison','Selected classical rules and two local models','Superiority to learned or non-myopic acquisition methods'],
        ['Behavioral evidence','Recorded choices and requested-ID checks','Causal effects of semantics, ordering or reasoning'],
        ['Resources and reliability','Sensor count, recorded time and process RSS','Energy, financial cost, failure recovery or operational safety']
    ],'Table 8. Evidence coverage and unresolved questions. Unmeasured properties are not inferred from acquisition validity or diagnostic accuracy.')
    sub('9.5 Future research directions')
    p('The first priority is independent confirmation under the frozen protocol, followed by the separately specified predictor and interface studies. These extensions should preserve untouched testing identities, report all planned contrasts and retain unfavorable outcomes. Closed-loop re-evaluation is essential: replaying old acquisition paths under a new classifier cannot show how a selector would respond to that classifier during acquisition. Randomized sensor ordering, consistent ID remapping, anonymized descriptions and score-visibility controls can help distinguish positional preferences from use of observed values or tool scores. Their results are not part of the present paper.')
    p('A stronger comparison should reproduce learned, amortized and non-myopic acquisition baselines under common splits, sensor information and validation budgets. Selection of the diagnostic predictor, probability calibration and acquisition hyperparameters should use training and validation data only. Repeating the resulting closed-loop comparison across independently fitted predictors and decoding seeds would assess whether policy rankings persist beyond the single fitted model studied here. Reporting both fit cost and per-episode computation would make the deployment trade-off explicit.')
    p('Temporal and external validation should then test early fault onset, changing process regimes, noisy or unavailable channels and measurement delays. Heterogeneous acquisition costs require a cost budget and an explicit stopping rule, together with accuracy–cost and accuracy–delay curves. Energy savings would require measured hardware power or sensor-level energy data; recorded inference time alone is insufficient. Evaluation on an unrelated industrial dataset and, where access permits, a non-intervening plant deployment would test transfer while preserving the separation between advisory measurement selection and safety-critical control.')
    h('10 Conclusion')
    if complete:
        p(f'This exploratory study combines 5,040 model episodes with paired accuracy estimates, runtime summaries, acquisition traces and supplementary predictor replay. At three sensors, assisted Qwen and Granite achieve {pct("qwen_assisted",3)} and {pct("granite_assisted",3)} balanced accuracy, below {pct("forward",3)} for validation-selected forward acquisition; the retained primary improvement criterion is not met. Three observations guide subsequent research: repeated calls need not produce adaptive trajectories, trajectory diversity need not improve diagnosis, and policy rankings can depend on the shared predictor. These observations support explicit tests of acquisition behavior alongside diagnostic performance and computation. Their scope is a single simulated snapshot setting, and causal explanations and transfer to other predictors or industrial systems remain open.')
    else:
        p('The completed classical pilot establishes a verified dataset conversion, a shared masked predictor, frozen acquisition references, and run-family uncertainty estimates. It does not yet answer whether repeated small language-model consultation earns its computational cost. The language-model execution and paired evaluation stages must be completed before that claim can be tested.')
    h('Data and code availability')
    p('The source data are available through Harvard Dataverse at https://doi.org/10.7910/DVN/6C3JR1 under the public-domain dedication and accompanying terms recorded in the authoritative metadata. The local research package contains code, source and model hashes, the analysis plan and amendments, split rules, predictions, action traces, timings and analysis outputs. A public code archive or DOI has not yet been created. No public code availability is claimed.')
    h('Declaration of generative AI assistance')
    p('OpenAI ChatGPT/Codex assisted with source checking, implementation, experimental execution, analysis scripting, data visualization and manuscript drafting. Qwen and Granite are experimental systems whose calls and outputs are logged separately. The human authors are responsible for verifying the scientific claims, references, code and manuscript and for approving the version submitted for publication.')
    h('References')
    refs=json.loads((ROOT/'paper/references.json').read_text())
    # Number citations by first occurrence and enforce a fully cited bibliography.
    pattern=re.compile(r'\[(\d+(?:,\s*\d+)*)\]')
    order=[]
    for match in pattern.finditer('\n'.join(markdown)):
        for number in map(int,match.group(1).split(',')):
            if number not in order:order.append(number)
    assert set(order)==set(range(1,len(refs)+1)), (order,len(refs))
    mapping={old:new for new,old in enumerate(order,1)}
    def renumber(text):
        return pattern.sub(lambda m:'['+', '.join(str(mapping[int(n)]) for n in m.group(1).split(','))+']',text)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if pattern.search(run.text):run.text=renumber(run.text)
    markdown=[renumber(text) for text in markdown]
    for number,old in enumerate(order,1):p(f'[{number}] '+refs[old-1].split('] ',1)[1])
    (ROOT/'paper/reference_numbering.json').write_text(json.dumps(mapping,indent=2))

    footer=sec.footer.paragraphs[0];footer.alignment=1
    footer.paragraph_format.space_after=Pt(0)
    footer.paragraph_format.line_spacing=1
    run=footer.add_run();run.font.name='Times New Roman';run.font.size=Pt(9)
    start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin');run._r.append(start)
    instruction=OxmlElement('w:instrText');instruction.text=' PAGE ';run._r.append(instruction)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.append(end)
    doc.core_properties.title=TITLE;doc.core_properties.author='';doc.core_properties.subject='Exploratory paired study with retained primary analysis'
    path=ROOT/'paper/pilot_manuscript.docx';doc.save(path)
    (ROOT/'paper/pilot_manuscript.md').write_text('\n'.join(markdown))
    if not (ROOT/'paper/author_information_needed.md').exists():
        (ROOT/'paper/author_information_needed.md').write_text('# Author information required before submission\n\nProvide the author list and order, affiliations, corresponding-author details, funding, competing interests, CRediT contributions, and approval of the AI-use declaration. The manuscript currently claims none of these on the authors’ behalf.\n')
    print(json.dumps(dict(docx=str(path),primary_pilot_complete=complete,abstract_words=len(abstract.split()))))
    subprocess.run([sys.executable,str(ROOT/'src/render_paper.py'),str(path),
                    '--output_dir',str(ROOT/'paper/rendered'),'--emit_pdf'],check=True)
    shutil.copy2(ROOT/'paper/rendered/pilot_manuscript.pdf',ROOT/'paper/manuscript.pdf')

if __name__=='__main__':main()
