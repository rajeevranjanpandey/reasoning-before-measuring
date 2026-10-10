"""Render the complete Markdown source into editable single-column Word."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def main():
    source = ROOT/'paper/manuscript.md'
    if not source.exists():
        raise RuntimeError('Integrate audited studies before Word authoring')
    lines = source.read_text().splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    for name, size in [('Normal',11),('Title',19),('Heading 1',13),('Heading 2',11)]:
        style = doc.styles[name]
        style.font.name='Times New Roman'
        style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+k),None)
        fonts.set(qn('w:ascii'),'Times New Roman')
        fonts.set(qn('w:hAnsi'),'Times New Roman')
    normal=doc.styles['Normal'].paragraph_format
    normal.space_after=Pt(6)
    normal.line_spacing=1.12
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.core_properties.title=lines[0].removeprefix('# ')
    def prose(text, style=None):
        p=doc.add_paragraph(style=style)
        pieces=text.split('x_S')
        for i,piece in enumerate(pieces):
            if i:
                run=p.add_run('x');run.italic=True
                run=p.add_run('S');run.font.subscript=True
            p.add_run(piece)
        if text.startswith(('Figure ','Table ')):
            p.paragraph_format.keep_together=True
        if re.match(r'^Table \d+\.',text):
            p.paragraph_format.keep_with_next=True
        return p
    def add_table(rows):
        n=len(rows[0])
        weights=[max(5,min(27,max(len(row[i]) for row in rows))) for i in range(n)]
        widths=[Inches(6.57*w/sum(weights)) for w in weights]
        t=doc.add_table(rows=1,cols=n)
        t.autofit=False
        for col,w in zip(t.columns,widths):col.width=w
        for cell,s in zip(t.rows[0].cells,rows[0]):cell.text=s
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement('w:'+edge)
            el.set(qn('w:val'),'single' if edge in ('top','bottom','insideH') else 'nil')
            el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        t._tbl.tblPr.append(borders)
        for row in rows[1:]:
            for cell,s in zip(t.add_row().cells,row):cell.text=s
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=width
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    for r in p.runs:r.font.size=Pt(9)
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[- :]+',v) for v in row):rows.append(row)
                i+=1
            assert all(len(row)==len(rows[0]) for row in rows)
            add_table(rows);continue
        image=re.fullmatch(r'!\[(.*?)\]\(([^)]+)\)',line)
        if image:
            p=doc.add_paragraph()
            p.add_run().add_picture(str(ROOT/'paper'/image[2]),width=Inches(6.3))
            j=i+1
            while j<len(lines) and not lines[j].strip():j+=1
            previous=next((value.strip() for value in reversed(lines[:i]) if value.strip()), '')
            caption_precedes=bool(re.match(r'^Figure \d+\.',previous))
            p.paragraph_format.keep_with_next=(not caption_precedes and j<len(lines) and bool(re.match(r'^Figure \d+\.',lines[j])))
            p.paragraph_format.keep_together=True
            for drawing in p._p.iter(qn('wp:docPr')):drawing.set('descr',image[1])
        elif line.startswith('### '):prose(line[4:],'Heading 2')
        elif line.startswith('## '):prose(line[3:],'Heading 1')
        elif line.startswith('# '):prose(line[2:],'Title')
        else:
            p=prose(line)
            # Existing pilot captions precede their image, new captions follow it.
            j=i+1
            while j<len(lines) and not lines[j].strip():j+=1
            if re.match(r'^Figure \d+\.',line) and j<len(lines) and lines[j].startswith('!['):
                p.paragraph_format.keep_with_next=True
        i+=1
    footer=section.footer.paragraphs[0]
    footer.alignment=1
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    out=ROOT/'paper/manuscript.docx'
    doc.save(out)
    print(out)


if __name__=='__main__':main()
